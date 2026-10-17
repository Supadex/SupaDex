#!/usr/bin/env python3
"""Create the retained SupaDAO technical-publication DOCX template."""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "publication" / "SupaDAO_Document_Template.docx"
PURPLE = RGBColor(166, 99, 229)
LIME = RGBColor(202, 255, 0)
NAVY = RGBColor(16, 24, 39)
MUTED = RGBColor(90, 103, 120)


def font(run, size, color=NAVY, bold=False, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size); run.font.color.rgb = color; run.bold = bold


def cell_shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("SUPADAO DEX  /  ")
    font(run, 8, MUTED)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(11); normal.font.color.rgb = NAVY
    normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.2
    for name, size, color, before, after in [
        ("Title", 30, NAVY, 0, 8), ("Subtitle", 14, MUTED, 0, 24),
        ("Heading 1", 18, PURPLE, 18, 9), ("Heading 2", 14, NAVY, 13, 6),
        ("Heading 3", 11, MUTED, 10, 4),
    ]:
        st = doc.styles[name]; st.font.name = "Arial"; st.font.size = Pt(size); st.font.color.rgb = color
        st.font.bold = name != "Subtitle"; st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)

    header = sec.header.paragraphs[0]
    header.text = "SUPADAO  /  TECHNICAL PUBLICATION"
    font(header.runs[0], 8, MUTED, True)
    add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(82); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SUPADAO"); font(r, 12, LIME, True)
    p = doc.add_paragraph("[DOCUMENT TITLE]", style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; remove_paragraph_border(p)
    p = doc.add_paragraph("[A precise subtitle describing the document's purpose]", style="Subtitle"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("VERSION [X.Y]  •  [DATE]  •  [STATUS]"); font(r, 9, MUTED, True)
    doc.add_paragraph("Use this publication system for evidence-led protocol reports. Replace bracketed fields, preserve the hierarchy, cite implementation sources, and distinguish verified behavior from design intent.")
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(38); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RISK NOTICE"); font(r, 9, PURPLE, True)
    p = doc.add_paragraph("This material is technical information, not legal, tax, or investment advice. State audit and deployment status explicitly."); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph("[Explain the system in plain English: what it is, the problem addressed, the architecture, current status, and the most important limitations.]")
    doc.add_heading("Key findings", level=2)
    for text in ["[Verified implementation fact with source reference]", "[Material limitation or residual risk]", "[Next validation milestone]"]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Technical section", level=1)
    doc.add_paragraph("[Describe contracts, data structures, calculations, execution paths, assumptions, and rounding behavior. Use monospace for identifiers and keep claims traceable to source.]")
    table = doc.add_table(rows=1, cols=3); table.autofit = False
    widths = [Inches(1.45), Inches(2.55), Inches(2.5)]
    for i, text in enumerate(["Item", "Evidence", "Status"]):
        table.columns[i].width = widths[i]; table.rows[0].cells[i].width = widths[i]
        table.rows[0].cells[i].text = text; cell_shade(table.rows[0].cells[i], "111827")
        font(table.rows[0].cells[i].paragraphs[0].runs[0], 9, RGBColor(255,255,255), True)
    row = table.add_row().cells
    for i, text in enumerate(["[Feature]", "[Contract/test/document]", "[Verified / partial / unverified]"]):
        row[i].width = widths[i]; row[i].text = text; font(row[i].paragraphs[0].runs[0], 9)

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("References and document control", level=1)
    doc.add_heading("References", level=2)
    doc.add_paragraph("1. [Primary source: title, author/organization, date, stable URL]", style="List Number")
    doc.add_paragraph("2. [Repository source: path, symbol, revision]", style="List Number")
    doc.add_heading("Revision record", level=2)
    doc.add_paragraph("Version [X.Y] — [Date] — [Status] — [Change summary]")
    doc.add_heading("Publication checklist", level=2)
    for item in ["Claims register reviewed", "Current tests recorded", "Deployment evidence checked", "PDF rendered and visually inspected", "No placeholder or unsupported claim remains"]:
        doc.add_paragraph("[ ] " + item)

    doc.core_properties.title = "SupaDAO Technical Publication Template"
    doc.core_properties.subject = "Reusable branded template for evidence-led SupaDAO protocol documents"
    OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
